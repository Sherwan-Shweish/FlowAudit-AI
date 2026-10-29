"""Tests für src/exporter.py – Excel-Prüfbericht (inkl. IKS-Hash-Kette) und Power-BI-Paket."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from src.demo import batchlauf
from src.exporter import (
    PROTOKOLL_SPALTEN,
    erstelle_excel_bericht,
    erstelle_powerbi_paket,
    lade_exportdaten,
    pruefe_hashkette,
)

from .conftest import STICHTAG


@pytest.fixture(scope="module")
def lauf(_demo_vorlage, tmp_path_factory):
    """Einmaliger Batchlauf (bucht alle 15 Musterdokumente in eine Kopie) – für alle Export-Tests."""
    from src.demo import stelle_demodaten_sicher

    ordner = tmp_path_factory.mktemp("export")
    import shutil

    shutil.copytree(_demo_vorlage.db_pfad.parent, ordner / "data")
    b = batchlauf(stelle_demodaten_sicher(ordner / "data"), Decimal("2.0"), ordner)
    daten = lade_exportdaten(b.db_kopie, STICHTAG)
    return b, daten, erstelle_excel_bericht(b.db_kopie, daten=daten), erstelle_powerbi_paket(b.db_kopie, daten=daten)


class TestExportdaten:
    def test_umfang(self, lauf):
        _, d, _, _ = lauf
        assert len(d.rechnungen) == 64 and len(d.protokoll) >= 14 * 8
        assert d.stichtag == STICHTAG

    def test_abweichungen_mit_wertmaessiger_auswirkung(self, lauf):
        _, d, _, _ = lauf
        nach_rechnung = {(a["XBLNR"], a["EBELP"]): a for a in d.abweichungen}
        preis = nach_rechnung[("TPK-2026-118", 10)]
        assert (preis["SOLL"], preis["IST"], preis["WERT_EUR"], preis["ZUSTAENDIG"]) == (4, Decimal("4.18"), 90, "Einkauf")
        teil = nach_rechnung[("RLV-2026-0447", 10)]
        assert (teil["MENGE_ABW"], teil["WERT_EUR"]) == (40, Decimal("596.00"))  # 40 Paletten × 14,90 €
        menge = nach_rechnung[("ODW-55160", 10)]
        assert (menge["MENGE_ABW"], menge["WERT_EUR"]) == (20, Decimal("138.00"))  # aus ZUGFeRD-PDF 14
        assert {a["PRUEFCODE"] for a in d.abweichungen} == {"PREISABWEICHUNG", "MENGENABWEICHUNG",
                                                            "KEIN_WARENEINGANG", "WARENEINGANG_UNVOLLSTAENDIG"}

    def test_skonto_risikoliste(self, lauf):
        _, d, _, _ = lauf
        erster = d.skonto[0]
        assert (erster["XBLNR"], erster["STATUS"], erster["RESTTAGE"], erster["BETRAG"]) == (
            "TPK-2026-118", "GEFÄHRDET", 2, Decimal("49.74"))
        status = [s["STATUS"] for s in d.skonto]
        assert status == sorted(status, key=["GEFÄHRDET", "VERLOREN", "GESICHERT"].index)
        summe = sum((s["BETRAG"] for s in d.skonto if s["STATUS"] == "GEFÄHRDET"), Decimal(0))
        assert summe == d.kpis.skonto_gefaehrdet_eur


class TestExcelBericht:
    def test_blaetter_tabellen_und_formeln(self, lauf):
        _, d, xlsx, _ = lauf
        wb = load_workbook(io.BytesIO(xlsx))
        assert wb.sheetnames == ["Übersicht", "Rechnungen", "3-Way-Abweichungen", "Skonto-Risiko", "IKS-Prüfprotokoll"]
        assert set(wb["Rechnungen"].tables) == {"tbl_Rechnungen"}
        assert set(wb["Skonto-Risiko"].tables) == {"tbl_Skonto"}
        assert "Stichtag" in wb.defined_names
        uebersicht = wb["Übersicht"]
        assert uebersicht["C9"].value.startswith("=IFERROR(COUNTIFS(tbl_Rechnungen[Workflow]")
        assert uebersicht["D9"].value.startswith("=IF(ABS(B9-C9)<0.01")
        assert uebersicht._charts, "Balkendiagramm der Sperrgründe fehlt"
        skonto = wb["Skonto-Risiko"]
        assert skonto["F2"].value == "=ROUND(D2*E2,2)" and skonto["H2"].value == "=G2-Stichtag"
        assert wb["Rechnungen"].max_row == 65  # Kopf + 64 Rechnungen

    def test_keine_bankdaten(self, lauf):
        _, d, xlsx, _ = lauf
        wb = load_workbook(io.BytesIO(xlsx))
        inhalt = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
        assert "IBAN" not in inhalt.replace("IBAN/SWIFT", "") and "DE14550501200200300400" not in inhalt

    def test_iks_protokoll_mit_hashes(self, lauf):
        _, d, xlsx, _ = lauf
        ws = load_workbook(io.BytesIO(xlsx))["IKS-Prüfprotokoll"]
        assert [c.value for c in ws[1]] == PROTOKOLL_SPALTEN
        zeile = [c.value for c in ws[2]]
        datei_sha, fakten_sha, kette = zeile[-3:]
        assert len(datei_sha) == len(fakten_sha) == len(kette) == 64
        assert ws.max_row - 1 == len(d.protokoll)

    def test_hashkette_intakt(self, lauf):
        assert pruefe_hashkette(lauf[2]) == (True, None)

    def test_hashkette_erkennt_manipulation(self, lauf):
        wb = load_workbook(io.BytesIO(lauf[2]))
        ws = wb["IKS-Prüfprotokoll"]
        ziel = next(r for r in ws.iter_rows(min_row=2) if r[5].value == "FEHLER")
        ziel[5].value = "OK"  # Befund nachträglich „weggeschönt“
        puffer = io.BytesIO()
        wb.save(puffer)
        intakt, ab = pruefe_hashkette(puffer.getvalue())
        assert not intakt and ab == ziel[0].value


class TestPowerBIPaket:
    def _zip(self, lauf) -> zipfile.ZipFile:
        return zipfile.ZipFile(io.BytesIO(lauf[3]))

    def test_inhalt(self, lauf):
        namen = {n.split("/", 1)[1] for n in self._zip(lauf).namelist()}
        assert namen == {"fact_rechnung.csv", "fact_pruefung.csv", "dim_lieferant.csv", "dim_workflow_schritt.csv",
                         "dim_pruefcode.csv", "dim_datum.csv", "measures.dax", "laden.pq", "README.txt", "manifest.json"}

    def _csv(self, lauf, name) -> list[list[str]]:
        roh = self._zip(lauf).read(f"flowaudit_powerbi/{name}")
        assert roh.startswith(b"\xef\xbb\xbf"), "UTF-8 mit BOM erwartet"
        return list(csv.reader(io.StringIO(roh.decode("utf-8-sig")), delimiter=";"))

    def test_deutsches_zahlen_und_datumsformat(self, lauf):
        zeilen = self._csv(lauf, "fact_rechnung.csv")
        kopf, daten = zeilen[0], zeilen[1:]
        assert len(daten) == 64
        zeile = next(z for z in daten if z[kopf.index("XBLNR")] == "TPK-2026-118")
        assert zeile[kopf.index("BRUTTO")] == "2487,10" and zeile[kopf.index("SKONTO_BETRAG")] == "49,74"
        assert zeile[kopf.index("SKONTO_STATUS")] == "GEFÄHRDET" and zeile[kopf.index("SKONTO_FRIST_ID")] == "20261030"
        assert zeile[kopf.index("EINGANG")].startswith("28.10.2026 ")
        assert all("." not in z[kopf.index("BRUTTO")] for z in daten), "kein Tausender-/Dezimalpunkt"

    def test_sternschema_referenzielle_integritaet(self, lauf):
        rechnung = self._csv(lauf, "fact_rechnung.csv")
        pruefung = self._csv(lauf, "fact_pruefung.csv")
        lieferant = self._csv(lauf, "dim_lieferant.csv")
        schritt = self._csv(lauf, "dim_workflow_schritt.csv")
        datum = self._csv(lauf, "dim_datum.csv")
        code = self._csv(lauf, "dim_pruefcode.csv")
        k = rechnung[0]
        ids = {z[0] for z in rechnung[1:]}
        assert {z[1] for z in pruefung[1:]} <= ids
        assert {z[k.index("LIFNR")] for z in rechnung[1:]} <= {z[0] for z in lieferant[1:]}
        assert {z[k.index("WORKFLOW_SCHRITT")] for z in rechnung[1:]} <= {z[0] for z in schritt[1:]}
        datum_ids = {z[0] for z in datum[1:]}
        for spalte in ("RECHNUNGSDATUM_ID", "EINGANG_DATUM_ID", "SKONTO_FRIST_ID"):
            assert {z[k.index(spalte)] for z in rechnung[1:] if z[k.index(spalte)]} <= datum_ids
        assert {z[2] for z in pruefung[1:]} <= {z[0] for z in code[1:]}
        assert "IBAN" not in lieferant[0] and "SWIFT" not in lieferant[0]

    def test_manifest_hashes(self, lauf):
        import hashlib

        z = self._zip(lauf)
        manifest = json.loads(z.read("flowaudit_powerbi/manifest.json"))
        assert manifest["stichtag"] == STICHTAG.isoformat()
        for name, info in manifest["dateien"].items():
            assert hashlib.sha256(z.read(f"flowaudit_powerbi/{name}")).hexdigest() == info["sha256"]
        assert manifest["dateien"]["fact_rechnung.csv"]["zeilen"] == 64

    def test_dax_und_power_query(self, lauf):
        z = self._zip(lauf)
        dax = z.read("flowaudit_powerbi/measures.dax").decode()
        assert "Dunkelverarbeitungsquote = DIVIDE" in dax
        assert '"de-DE"' in z.read("flowaudit_powerbi/laden.pq").decode()
