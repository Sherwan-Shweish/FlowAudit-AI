"""
FlowAudit-AI · Export-Modul: Excel-Prüfbericht & Power-BI-Paket
===============================================================

1. **Excel-Prüfbericht** (``.xlsx``) – fünf Blätter:

   * ``Übersicht`` – Stichtag, KPIs (Python) **und** Gegenrechnung mit Excel-Formeln
     (``COUNTIFS``/``SUMIFS`` auf strukturierte Tabellenverweise), Sperrgründe mit Balkendiagramm,
   * ``Rechnungen`` – Rechnungsübersicht als Excel-Tabelle mit Ampel-Formatierung und
     Durchlaufzeit als Formel,
   * ``3-Way-Abweichungen`` – Preis-, Mengen- und Wareneingangsabweichungen mit Soll/Ist,
     EKPO-Bezug und wertmäßiger Auswirkung (Datenbalken),
   * ``Skonto-Risiko`` – Skontobetrag und Resttage als Formeln (benannter Bereich ``Stichtag``),
     Status gesichert/gefährdet/verloren,
   * ``IKS-Prüfprotokoll`` – jede Einzelprüfung inkl. SHA-256 der Eingangsdatei, SHA-256 des
     Prüfstands und einer **Hash-Kette** über alle Zeilen (nachträgliche Änderungen im Export
     werden von :func:`pruefe_hashkette` erkannt).

2. **Power-BI-Paket** (``.zip``) – Sternschema als CSV (UTF-8 mit BOM, Semikolon, Dezimalkomma,
   Datum TT.MM.JJJJ) mit Faktentabellen, Dimensionen inkl. Kalendertabelle, DAX-Measures,
   Power-Query-Skript (Kultur ``de-DE``) und Manifest mit SHA-256 je Datei.

Bankdaten (``LFBK``: IBAN/SWIFT) werden grundsätzlich nicht exportiert.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Iterable, Optional

from .agent import EMPFAENGER_JE_SCHRITT, URTEIL_JE_SCHRITT, PruefCode, WorkflowSchritt
from .bpmn_tracker import ProzessKPIs, berechne_kpis, lade_vorgaenge
from .matching_engine import lese_stichtag

__all__ = [
    "ExportDaten",
    "lade_exportdaten",
    "erstelle_excel_bericht",
    "erstelle_powerbi_paket",
    "speichere_excel_bericht",
    "speichere_powerbi_paket",
    "pruefe_hashkette",
    "HASH_START",
]

HASH_START = "0" * 64
ABWEICHUNGS_CODES = (PruefCode.PREISABWEICHUNG, PruefCode.MENGENABWEICHUNG,
                     PruefCode.KEIN_WARENEINGANG, PruefCode.WARENEINGANG_UNVOLLSTAENDIG)
SCHRITT_TEXT = {
    "TOUCHLESS_FREIGABE": "Touchless-Freigabe",
    "SPERRE_DUBLETTE": "Sperre: Dublette",
    "KLAERUNG_LIEFERANT_FORMFEHLER": "Klärung Lieferant: Formfehler",
    "COMPLIANCE_MAVERICK_BUYING": "Compliance: Maverick Buying",
    "SPERRE_FEHLENDER_WARENEINGANG": "Sperre: Wareneingang",
    "KLAERUNG_EINKAUF_PREIS": "Klärung Einkauf: Preis",
    "KLAERUNG_EINKAUF_MENGE": "Klärung Einkauf: Menge",
}
PRUEFCODE_INFO = {  # Code → (Bezeichnung, Kategorie, BPMN-Gateway)
    "DUBLETTE": ("Doppelte Rechnung", "Dublettenprüfung", "Dublettenprüfung"),
    "PFLICHTFELD_FEHLT": ("Pflichtangabe § 14 Abs. 4 UStG fehlt", "Formal", "Formale Prüfung"),
    "USTID_UNGUELTIG": ("USt-IdNr. ungültig", "Formal", "Formale Prüfung"),
    "LEITWEG_ID_FEHLT": ("Käuferreferenz/Leitweg-ID (BT-10) fehlt", "Formal", "Formale Prüfung"),
    "STEUERSATZ_UNZULAESSIG": ("Unzulässiger Steuersatz", "Formal", "Formale Prüfung"),
    "SUMMENFEHLER": ("Rechenfehler", "Formal", "Formale Prüfung"),
    "KEINE_BESTELLUNG": ("Keine gültige Bestellung (Maverick Buying)", "Compliance", "Bestellbezug"),
    "KEIN_WARENEINGANG": ("Kein Wareneingang gebucht", "3-Way-Matching", "3-Way-Abgleich"),
    "WARENEINGANG_UNVOLLSTAENDIG": ("Wareneingang unvollständig", "3-Way-Matching", "3-Way-Abgleich"),
    "PREISABWEICHUNG": ("Preisabweichung außerhalb Toleranz", "3-Way-Matching", "3-Way-Abgleich"),
    "MENGENABWEICHUNG": ("Menge über offener Bestellmenge", "3-Way-Matching", "3-Way-Abgleich"),
    "SKONTO_FRIST": ("Skontofrist läuft ab / abgelaufen", "Hinweis", "–"),
    "VERDAECHTIGER_FREITEXT": ("Verdacht auf Prompt-Injection", "Hinweis", "–"),
}


def _r2(x: Decimal) -> Decimal:
    return x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _dez(x: Any) -> Optional[Decimal]:
    return None if x is None else Decimal(str(round(float(x), 6)))


def _dt(text: Optional[str]) -> Optional[datetime]:
    return datetime.fromisoformat(text) if text else None


def _d(text: Optional[str]) -> Optional[date]:
    return date.fromisoformat(text[:10]) if text else None


# ═════════════════════════════════════════════════════════════════════════════
# 1 · Daten laden (schreibgeschützt) und aufbereiten
# ═════════════════════════════════════════════════════════════════════════════

@dataclass
class ExportDaten:
    stichtag: date
    kpis: ProzessKPIs
    rechnungen: list[dict] = field(default_factory=list)
    abweichungen: list[dict] = field(default_factory=list)
    skonto: list[dict] = field(default_factory=list)
    protokoll: list[dict] = field(default_factory=list)
    lieferanten: list[dict] = field(default_factory=list)


def _abfrage(con: sqlite3.Connection, sql: str, *param: Any) -> list[dict]:
    cur = con.execute(sql, param)
    spalten = [d[0] for d in cur.description]
    return [dict(zip(spalten, z)) for z in cur.fetchall()]


def _skonto_status(abschluss: Optional[datetime], frist: date, stichtag: date) -> str:
    if abschluss is not None:
        return "GESICHERT" if abschluss.date() <= frist else "VERLOREN"
    return "VERLOREN" if stichtag > frist else "GEFÄHRDET"


def lade_exportdaten(db_pfad: Path | str, stichtag: Optional[date] = None) -> ExportDaten:
    db_pfad = Path(db_pfad)
    stichtag = stichtag or lese_stichtag(db_pfad) or date.today()
    con = sqlite3.connect(f"{db_pfad.resolve().as_uri()}?mode=ro", uri=True)
    try:
        rechnungen = _abfrage(con, """
            SELECT k.BELNR, k.GJAHR, k.XBLNR, k.LIFNR, l.NAME1, k.BLDAT, k.ZZ_EBELN, k.ZZ_FORMAT,
                   k.ZZ_NETTO, k.WMWST1, k.RMWWR, k.RBSTAT, k.ZLSPR, k.ZZ_WF_SCHRITT, k.ZZ_URTEIL,
                   k.ZZ_EINGANG, k.ZZ_ABSCHLUSS, k.ZZ_SKONTO_PROZ, k.ZZ_SKONTO_FRIST, k.ZZ_SHA256
            FROM RBKP k LEFT JOIN LFA1 l ON l.LIFNR = k.LIFNR
            WHERE k.STBLG IS NULL AND k.ZZ_WF_SCHRITT IS NOT NULL
            ORDER BY k.ZZ_EINGANG, k.BELNR""")
        protokoll = _abfrage(con, """
            SELECT p.BELNR, p.GJAHR, p.XBLNR, p.PRUEFCODE, p.STATUS, p.MELDUNG, p.SOLL, p.IST, p.ABWEICHUNG,
                   p.ABWEICHUNG_PROZ, p.REFERENZ, p.GEPRUEFT_AM, k.ZZ_SHA256 AS DATEI_SHA256, p.FAKTEN_SHA256
            FROM ZFA_PRUEFPROT p LEFT JOIN RBKP k ON k.BELNR = p.BELNR AND k.GJAHR = p.GJAHR
            ORDER BY p.BELNR, p.GJAHR, p.rowid""")
        ekpo = {(z["EBELN"], z["EBELP"]): z for z in _abfrage(con, "SELECT EBELN, EBELP, TXZ01, NETPR, PEINH FROM EKPO")}
        lieferanten = _abfrage(con, "SELECT LIFNR, NAME1, STRAS, PSTLZ, ORT01, LAND1, STCEG, ZTERM FROM LFA1 ORDER BY LIFNR")
    finally:
        con.close()

    rechnung_von = {(r["BELNR"], r["GJAHR"]): r for r in rechnungen}

    abweichungen = []
    for p in protokoll:
        if p["STATUS"] != "FEHLER" or p["PRUEFCODE"] not in {c.value for c in ABWEICHUNGS_CODES}:
            continue
        r = rechnung_von.get((p["BELNR"], p["GJAHR"]), {})
        m = re.search(r"EKPO (\S+)/(\d{5})", p["REFERENZ"] or "")
        pos = ekpo.get((m.group(1), int(m.group(2)))) if m else None
        preis = _dez(pos["NETPR"]) / Decimal(pos["PEINH"] or 1) if pos else None
        soll, ist, abw = _dez(p["SOLL"]), _dez(p["IST"]), _dez(p["ABWEICHUNG"])
        if p["PRUEFCODE"] == PruefCode.PREISABWEICHUNG.value:
            menge_abw, wert = None, abw                       # Abweichung ist bereits die €-Auswirkung
        else:
            menge_abw = abw if abw is not None else (ist - soll if ist is not None and soll is not None else None)
            wert = _r2(menge_abw * preis) if menge_abw is not None and preis is not None else None
        schritt = r.get("ZZ_WF_SCHRITT")
        abweichungen.append({
            "BELNR": p["BELNR"], "XBLNR": p["XBLNR"], "LIEFERANT": r.get("NAME1"), "EBELN": m.group(1) if m else r.get("ZZ_EBELN"),
            "EBELP": int(m.group(2)) if m else None, "TEXT": pos["TXZ01"] if pos else None, "PRUEFCODE": p["PRUEFCODE"],
            "SOLL": soll, "IST": ist, "MENGE_ABW": menge_abw, "ABW_PROZ": _dez(p["ABWEICHUNG_PROZ"]), "WERT_EUR": wert,
            "WORKFLOW": schritt, "ZUSTAENDIG": (EMPFAENGER_JE_SCHRITT[WorkflowSchritt(schritt)].value.title()
                                               if schritt and WorkflowSchritt(schritt) in EMPFAENGER_JE_SCHRITT else "–"),
        })

    skonto = []
    for r in rechnungen:
        if r["ZZ_SKONTO_PROZ"] is None or not r["ZZ_SKONTO_FRIST"]:
            continue
        frist = _d(r["ZZ_SKONTO_FRIST"])
        status = _skonto_status(_dt(r["ZZ_ABSCHLUSS"]), frist, stichtag)
        skonto.append({
            "BELNR": r["BELNR"], "XBLNR": r["XBLNR"], "LIEFERANT": r["NAME1"], "BRUTTO": _dez(r["RMWWR"]),
            "PROZ": _dez(r["ZZ_SKONTO_PROZ"]), "BETRAG": _r2(_dez(r["RMWWR"]) * _dez(r["ZZ_SKONTO_PROZ"]) / 100),
            "FRIST": frist, "RESTTAGE": (frist - stichtag).days, "STATUS": status,
            "WORKFLOW": r["ZZ_WF_SCHRITT"], "ABGESCHLOSSEN": _dt(r["ZZ_ABSCHLUSS"]),
        })
    rang = {"GEFÄHRDET": 0, "VERLOREN": 1, "GESICHERT": 2}
    skonto.sort(key=lambda s: (rang[s["STATUS"]], s["RESTTAGE"], -s["BETRAG"]))

    kpis = berechne_kpis(lade_vorgaenge(db_pfad), stichtag)
    return ExportDaten(stichtag, kpis, rechnungen, abweichungen, skonto, protokoll, lieferanten)


# ═════════════════════════════════════════════════════════════════════════════
# 2 · Hash-Kette für das IKS-Prüfprotokoll
# ═════════════════════════════════════════════════════════════════════════════

def _zeilen_hash(vorgaenger: str, werte: Iterable[Any]) -> str:
    # Zahlen einheitlich als float: Excel liest 4.0 als 4 zurück – die Kette muss das aushalten.
    kanonisch = json.dumps([w if w is None or isinstance(w, str) else float(w) for w in werte],
                           ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256((vorgaenger + kanonisch).encode("utf-8")).hexdigest()


PROTOKOLL_SPALTEN = ["Lfd. Nr.", "BELNR", "GJAHR", "Rechnungsnr.", "Prüfcode", "Status", "Meldung", "Soll", "Ist",
                     "Abweichung", "Abweichung %", "Referenz", "Geprüft am", "SHA-256 Eingangsdatei",
                     "SHA-256 Prüfstand", "Zeilen-Hash (Kette)"]


def _protokoll_zeilen(protokoll: list[dict]) -> list[list[Any]]:
    zeilen, vorgaenger = [], HASH_START
    for i, p in enumerate(protokoll, start=1):
        werte = [i, p["BELNR"], p["GJAHR"], p["XBLNR"], p["PRUEFCODE"], p["STATUS"], p["MELDUNG"],
                 p["SOLL"], p["IST"], p["ABWEICHUNG"], p["ABWEICHUNG_PROZ"], p["REFERENZ"], p["GEPRUEFT_AM"],
                 p["DATEI_SHA256"], p["FAKTEN_SHA256"]]
        vorgaenger = _zeilen_hash(vorgaenger, werte)
        zeilen.append(werte + [vorgaenger])
    return zeilen


def pruefe_hashkette(xlsx: bytes | Path | str) -> tuple[bool, Optional[int]]:
    """Prüft die Hash-Kette im Blatt ``IKS-Prüfprotokoll``. Rückgabe: (intakt?, erste fehlerhafte Lfd. Nr.)."""
    from openpyxl import load_workbook

    quelle = io.BytesIO(xlsx) if isinstance(xlsx, bytes) else xlsx
    ws = load_workbook(quelle, data_only=True)["IKS-Prüfprotokoll"]
    vorgaenger = HASH_START
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        if zeile[0] is None:
            continue
        *werte, gespeichert = zeile[: len(PROTOKOLL_SPALTEN)]
        vorgaenger = _zeilen_hash(vorgaenger, werte)
        if vorgaenger != gespeichert:
            return False, int(zeile[0])
    return True, None


# ═════════════════════════════════════════════════════════════════════════════
# 3 · Excel-Prüfbericht
# ═════════════════════════════════════════════════════════════════════════════

EUR = '#,##0.00 "€";[Red]-#,##0.00 "€"'
ZAHL = "#,##0.00"
MENGE = "#,##0.###"
DATUM = "DD.MM.YYYY"
ZEIT = "DD.MM.YYYY HH:MM"


def erstelle_excel_bericht(db_pfad: Path | str, stichtag: Optional[date] = None,
                           daten: Optional[ExportDaten] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table, TableStyleInfo

    d = daten or lade_exportdaten(db_pfad, stichtag)
    wb = Workbook()
    blau, grau = "1D4ED8", "57606A"
    gruen_f, gelb_f, rot_f = (PatternFill("solid", fgColor=c) for c in ("D1F2D9", "FFF4CC", "FDE2E1"))
    duenn = Border(bottom=Side(style="thin", color="D0D7DE"))

    def tabelle(ws, name: str, spalten: list[str], zeilen: list[list[Any]], formate: dict[str, str],
                breiten: dict[str, int], start: int = 1, stil: str = "TableStyleMedium2") -> str:
        """Schreibt eine Excel-Tabelle (ListObject) und liefert den Bereich."""
        for j, s in enumerate(spalten, start=1):
            ws.cell(start, j, s)
        for i, z in enumerate(zeilen, start=start + 1):
            for j, w in enumerate(z, start=1):
                ws.cell(i, j, float(w) if isinstance(w, Decimal) else w)
        ende = start + max(len(zeilen), 1)
        bereich = f"A{start}:{get_column_letter(len(spalten))}{ende}"
        t = Table(displayName=name, ref=bereich)
        t.tableStyleInfo = TableStyleInfo(name=stil, showRowStripes=True)
        ws.add_table(t)
        for j, s in enumerate(spalten, start=1):
            buchstabe = get_column_letter(j)
            ws.column_dimensions[buchstabe].width = breiten.get(s, max(10, len(s) + 3))
            if s in formate:
                for i in range(start + 1, ende + 1):
                    ws[f"{buchstabe}{i}"].number_format = formate[s]
        ws.freeze_panes = ws.cell(start + 1, 1)
        return bereich

    def spalte(spalten: list[str], name: str) -> str:
        return get_column_letter(spalten.index(name) + 1)

    # ── Übersicht ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Übersicht"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "FlowAudit-AI · Prüfbericht Purchase-to-Pay"
    ws["A1"].font = Font(size=16, bold=True, color=blau)
    ws["A2"] = "Synthetische Demo-Daten – RheinMain Industrietechnik GmbH (fiktiv). Keine Steuer- oder Rechtsberatung."
    ws["A2"].font = Font(italic=True, color=grau)
    ws["A4"], ws["B4"] = "Stichtag", d.stichtag
    ws["B4"].number_format = DATUM
    ws["A5"], ws["B5"] = "Erstellt am", datetime.now().replace(microsecond=0)
    ws["B5"].number_format = ZEIT
    ws["A6"], ws["B6"] = "Rechnungen im Bestand", len(d.rechnungen)
    wb.defined_names["Stichtag"] = DefinedName("Stichtag", attr_text="'Übersicht'!$B$4")

    k = d.kpis
    kpi_zeilen = [
        ("Dunkelverarbeitungsquote", float(k.dunkelverarbeitungsquote_prozent) / 100, "0.00%",
         '=IFERROR(COUNTIFS(tbl_Rechnungen[Workflow],"TOUCHLESS_FREIGABE")/ROWS(tbl_Rechnungen[Workflow]),0)'),
        ("Maverick-Buying-Quote", float(k.maverick_quote_prozent) / 100, "0.00%",
         '=IFERROR(COUNTIFS(tbl_Rechnungen[Workflow],"COMPLIANCE_MAVERICK_BUYING")/ROWS(tbl_Rechnungen[Workflow]),0)'),
        ("Offene Vorgänge", k.anzahl_offen, "0", '=COUNTBLANK(tbl_Rechnungen[Abschluss])'),
        ("Offenes Volumen (brutto)", float(k.volumen_offen_eur), EUR,
         '=SUMIFS(tbl_Rechnungen[Brutto],tbl_Rechnungen[Abschluss],"")'),
        ("Ø Durchlaufzeit gesamt (h)", float(k.durchlaufzeit_avg_stunden or 0), ZAHL,
         '=IFERROR(AVERAGE(tbl_Rechnungen[Durchlaufzeit h]),0)'),
        ("Ø Durchlaufzeit Touchless (min)", float(k.durchlaufzeit_touchless_avg_minuten or 0), ZAHL,
         '=IFERROR(AVERAGEIFS(tbl_Rechnungen[Durchlaufzeit h],tbl_Rechnungen[Workflow],"TOUCHLESS_FREIGABE")*60,0)'),
        ("Skonto gesichert", float(k.skonto_gesichert_eur), EUR, '=SUMIFS(tbl_Skonto[Skontobetrag],tbl_Skonto[Status],"GESICHERT")'),
        ("Skonto gefährdet", float(k.skonto_gefaehrdet_eur), EUR, '=SUMIFS(tbl_Skonto[Skontobetrag],tbl_Skonto[Status],"GEFÄHRDET")'),
        ("Skonto verloren", float(k.skonto_verloren_eur), EUR, '=SUMIFS(tbl_Skonto[Skontobetrag],tbl_Skonto[Status],"VERLOREN")'),
        ("Wertmäßige 3-Way-Abweichungen", float(sum((a["WERT_EUR"] or 0 for a in d.abweichungen), Decimal(0))), EUR,
         "=SUM(tbl_Abweichungen[Auswirkung EUR])"),
    ]
    ws["A8"], ws["B8"], ws["C8"], ws["D8"] = "Kennzahl", "Wert (Python-KPI-Engine)", "Gegenrechnung (Excel-Formel)", "Abgleich"
    for zelle in ws[8]:
        zelle.font, zelle.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=blau)
    for i, (name, wert, fmt, formel) in enumerate(kpi_zeilen, start=9):
        ws.cell(i, 1, name).border = duenn
        ws.cell(i, 2, wert).number_format = fmt
        ws.cell(i, 3, formel).number_format = fmt
        ws.cell(i, 4, f'=IF(ABS(B{i}-C{i})<0.01,"✔ stimmt","✘ prüfen")')
        for j in (2, 3, 4):
            ws.cell(i, j).border = duenn
    letzte = 8 + len(kpi_zeilen)
    ws.conditional_formatting.add(f"D9:D{letzte}", FormulaRule(formula=['LEFT(D9,1)="✔"'], fill=gruen_f))
    ws.conditional_formatting.add(f"D9:D{letzte}", FormulaRule(formula=['LEFT(D9,1)="✘"'], fill=rot_f))

    s0 = letzte + 2
    ws.cell(s0, 1, "Sperr- und Klärungsgründe").font = Font(bold=True, color=blau, size=12)
    ws.cell(s0 + 1, 1, "Grund")
    ws.cell(s0 + 1, 2, "Anzahl")
    ws.cell(s0 + 1, 3, "Anteil")
    for zelle in ws[s0 + 1][:3]:
        zelle.font, zelle.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=blau)
    for i, (schritt, anzahl) in enumerate(k.sperrgruende.items(), start=s0 + 2):
        ws.cell(i, 1, SCHRITT_TEXT.get(schritt, schritt))
        ws.cell(i, 2, f'=COUNTIFS(tbl_Rechnungen[Workflow],"{schritt}")')
        ws.cell(i, 3, f"=IFERROR(B{i}/SUM($B${s0 + 2}:$B${s0 + 1 + len(k.sperrgruende)}),0)").number_format = "0.0%"
    if k.sperrgruende:
        diagramm = BarChart()
        diagramm.type, diagramm.style = "bar", 10
        diagramm.title = "Sperr- und Klärungsgründe"
        diagramm.y_axis.title, diagramm.legend = "Anzahl", None
        n = len(k.sperrgruende)
        diagramm.add_data(Reference(ws, min_col=2, min_row=s0 + 1, max_row=s0 + 1 + n), titles_from_data=True)
        diagramm.set_categories(Reference(ws, min_col=1, min_row=s0 + 2, max_row=s0 + 1 + n))
        diagramm.height, diagramm.width = 7.5, 16
        ws.add_chart(diagramm, f"F{s0}")
    hinweise = [
        "Lesehilfe:",
        "• Spalte B stammt aus der Python-KPI-Engine (src/bpmn_tracker.py), Spalte C rechnet dieselbe Kennzahl mit Excel-Formeln",
        "  auf den Tabellen nach – beide Wege müssen übereinstimmen (Spalte D).",
        "• IKS-Prüfprotokoll: Jede Zeile enthält einen SHA-256-Hash aus ihrem Inhalt und dem Hash der Vorzeile (Hash-Kette).",
        "  Wird eine Zeile nachträglich verändert, bricht die Kette ab diesem Punkt – prüfbar mit src.exporter.pruefe_hashkette().",
        "• Bankdaten (LFBK: IBAN/SWIFT) sind nicht Bestandteil dieses Berichts.",
    ]
    for i, text in enumerate(hinweise, start=s0 + 3 + len(k.sperrgruende)):
        ws.cell(i, 1, text).font = Font(bold=text.endswith(":"), color=grau)
    for b, w in zip("ABCD", (34, 26, 30, 12)):
        ws.column_dimensions[b].width = w

    # ── Rechnungen ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("Rechnungen")
    sp = ["BELNR", "GJAHR", "Rechnungsnr.", "LIFNR", "Lieferant", "Rechnungsdatum", "Bestellung", "Format",
          "Netto", "Steuer", "Brutto", "Workflow", "Urteil", "Zahlsperre", "Eingang", "Abschluss",
          "Durchlaufzeit h", "SHA-256 Eingangsdatei"]
    zeilen = []
    for i, r in enumerate(d.rechnungen, start=2):
        zeilen.append([r["BELNR"], r["GJAHR"], r["XBLNR"], r["LIFNR"], r["NAME1"], _d(r["BLDAT"]), r["ZZ_EBELN"],
                       r["ZZ_FORMAT"], _dez(r["ZZ_NETTO"]), _dez(r["WMWST1"]), _dez(r["RMWWR"]), r["ZZ_WF_SCHRITT"],
                       r["ZZ_URTEIL"], r["ZLSPR"], _dt(r["ZZ_EINGANG"]), _dt(r["ZZ_ABSCHLUSS"]),
                       f'=IF(P{i}="","",(P{i}-O{i})*24)', r["ZZ_SHA256"] or "– (Altbestand)"])
    tabelle(ws, "tbl_Rechnungen", sp, zeilen,
            {"Netto": EUR, "Steuer": EUR, "Brutto": EUR, "Rechnungsdatum": DATUM, "Eingang": ZEIT, "Abschluss": ZEIT,
             "Durchlaufzeit h": ZAHL},
            {"Lieferant": 34, "Workflow": 32, "Rechnungsnr.": 16, "Eingang": 17, "Abschluss": 17,
             "SHA-256 Eingangsdatei": 30, "BELNR": 12, "Bestellung": 12, "Format": 15})
    u = spalte(sp, "Urteil")
    ende = len(zeilen) + 1
    for wert, fuellung in (("FREIGABE", gruen_f), ("KLAERUNG", gelb_f), ("SPERRE", rot_f)):
        ws.conditional_formatting.add(f"{u}2:{u}{ende}", CellIsRule(operator="equal", formula=[f'"{wert}"'], fill=fuellung))

    # ── 3-Way-Abweichungen ─────────────────────────────────────────────────
    ws = wb.create_sheet("3-Way-Abweichungen")
    sp = ["BELNR", "Rechnungsnr.", "Lieferant", "Bestellung", "EBELP", "Position (EKPO)", "Prüfcode", "Soll", "Ist",
          "Abweichung Menge", "Abweichung %", "Auswirkung EUR", "Workflow", "Zuständig"]
    zeilen = [[a["BELNR"], a["XBLNR"], a["LIEFERANT"], a["EBELN"], a["EBELP"], a["TEXT"], a["PRUEFCODE"], a["SOLL"],
               a["IST"], a["MENGE_ABW"], None if a["ABW_PROZ"] is None else a["ABW_PROZ"] / 100, a["WERT_EUR"],
               a["WORKFLOW"], a["ZUSTAENDIG"]] for a in d.abweichungen]
    tabelle(ws, "tbl_Abweichungen", sp, zeilen, {"Soll": MENGE, "Ist": MENGE, "Abweichung Menge": MENGE,
                                                  "Abweichung %": "0.00%", "Auswirkung EUR": EUR},
            {"Lieferant": 32, "Position (EKPO)": 30, "Prüfcode": 30, "Workflow": 32, "Rechnungsnr.": 16,
             "Bestellung": 12, "BELNR": 12}, stil="TableStyleMedium3")
    if zeilen:
        w = spalte(sp, "Auswirkung EUR")
        ws.conditional_formatting.add(f"{w}2:{w}{len(zeilen) + 1}",
                                      DataBarRule(start_type="num", start_value=0, end_type="max", color="CF222E"))
    n = len(zeilen) + 3
    ws.cell(n, 11, "Summe").font = Font(bold=True)
    ws.cell(n, 12, "=SUBTOTAL(109,tbl_Abweichungen[Auswirkung EUR])").number_format = EUR
    ws.cell(n, 12).font = Font(bold=True)
    ws.cell(n + 1, 1, "Preisabweichung: Soll/Ist = Einzelpreis, Auswirkung = (Ist − Soll) × Menge. "
                      "Mengen-/WE-Abweichung: Soll/Ist = Menge, Auswirkung = Mehrmenge × Bestellpreis (EKPO).").font = Font(
        italic=True, color=grau)

    # ── Skonto-Risiko ──────────────────────────────────────────────────────
    ws = wb.create_sheet("Skonto-Risiko")
    sp = ["BELNR", "Rechnungsnr.", "Lieferant", "Brutto", "Skonto %", "Skontobetrag", "Skontofrist", "Resttage",
          "Status", "Workflow", "Abgeschlossen"]
    zeilen = []
    for i, s in enumerate(d.skonto, start=2):
        zeilen.append([s["BELNR"], s["XBLNR"], s["LIEFERANT"], s["BRUTTO"], s["PROZ"] / 100,
                       f"=ROUND(D{i}*E{i},2)", s["FRIST"], f"=G{i}-Stichtag", s["STATUS"], s["WORKFLOW"],
                       s["ABGESCHLOSSEN"]])
    tabelle(ws, "tbl_Skonto", sp, zeilen, {"Brutto": EUR, "Skonto %": "0.00%", "Skontobetrag": EUR,
                                            "Skontofrist": DATUM, "Resttage": "0", "Abgeschlossen": ZEIT},
            {"Lieferant": 34, "Workflow": 32, "Rechnungsnr.": 16, "Abgeschlossen": 17, "BELNR": 12},
            stil="TableStyleMedium7")
    if zeilen:
        st_, rt = spalte(sp, "Status"), spalte(sp, "Resttage")
        e = len(zeilen) + 1
        ws.conditional_formatting.add(f"{st_}2:{st_}{e}", CellIsRule(operator="equal", formula=['"GESICHERT"'], fill=gruen_f))
        ws.conditional_formatting.add(f"{st_}2:{st_}{e}", CellIsRule(operator="equal", formula=['"GEFÄHRDET"'], fill=gelb_f))
        ws.conditional_formatting.add(f"{st_}2:{st_}{e}", CellIsRule(operator="equal", formula=['"VERLOREN"'], fill=rot_f))
        ws.conditional_formatting.add(f"{rt}2:{rt}{e}", FormulaRule(
            formula=[f'AND(${st_}2="GEFÄHRDET",{rt}2<=3)'], fill=rot_f, font=Font(bold=True, color="CF222E")))
    n = len(zeilen) + 3
    ws.cell(n, 5, "Summe gefährdet").font = Font(bold=True)
    ws.cell(n, 6, '=SUMIFS(tbl_Skonto[Skontobetrag],tbl_Skonto[Status],"GEFÄHRDET")').number_format = EUR
    ws.cell(n + 1, 1, "Resttage = Skontofrist − Stichtag (benannter Bereich „Stichtag“ auf dem Blatt Übersicht). "
                      "Rot: gefährdeter Skonto mit höchstens 3 Resttagen.").font = Font(italic=True, color=grau)

    # ── IKS-Prüfprotokoll ──────────────────────────────────────────────────
    ws = wb.create_sheet("IKS-Prüfprotokoll")
    zeilen = _protokoll_zeilen(d.protokoll)
    tabelle(ws, "tbl_Pruefprotokoll", PROTOKOLL_SPALTEN, zeilen,
            {"Soll": "General", "Ist": "General", "Abweichung": "General", "Abweichung %": "General"},
            {"Meldung": 60, "Prüfcode": 28, "Referenz": 26, "Geprüft am": 20, "SHA-256 Eingangsdatei": 66,
             "SHA-256 Prüfstand": 66, "Zeilen-Hash (Kette)": 66, "Rechnungsnr.": 16, "BELNR": 12}, stil="TableStyleLight9")
    if zeilen:
        s = get_column_letter(PROTOKOLL_SPALTEN.index("Status") + 1)
        e = len(zeilen) + 1
        ws.conditional_formatting.add(f"{s}2:{s}{e}", CellIsRule(operator="equal", formula=['"FEHLER"'], fill=rot_f))
        ws.conditional_formatting.add(f"{s}2:{s}{e}", CellIsRule(operator="equal", formula=['"WARNUNG"'], fill=gelb_f))
    for z in ws.iter_rows(min_row=2, max_col=len(PROTOKOLL_SPALTEN)):
        z[PROTOKOLL_SPALTEN.index("Meldung")].alignment = Alignment(wrap_text=False)

    for blatt in wb.worksheets:  # Druck: Querformat, auf eine Seitenbreite
        blatt.page_setup.orientation = "landscape"
        blatt.page_setup.paperSize = blatt.PAPERSIZE_A4
        blatt.page_setup.fitToWidth, blatt.page_setup.fitToHeight = 1, 0
        blatt.sheet_properties.pageSetUpPr.fitToPage = True
        blatt.oddFooter.center.text = "FlowAudit-AI · &A · Seite &P von &N"
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = "FlowAudit-AI"
    wb.properties.title = "FlowAudit-AI Prüfbericht (synthetische Demo-Daten)"
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


# ═════════════════════════════════════════════════════════════════════════════
# 4 · Power-BI-Paket (Sternschema)
# ═════════════════════════════════════════════════════════════════════════════

def _de(wert: Any, stellen: Optional[int] = None) -> str:
    """Deutsches CSV-Format: Dezimalkomma, kein Tausendertrennzeichen, Datum TT.MM.JJJJ."""
    if wert is None:
        return ""
    if isinstance(wert, bool):
        return "1" if wert else "0"
    if isinstance(wert, datetime):
        return wert.strftime("%d.%m.%Y %H:%M:%S")
    if isinstance(wert, date):
        return wert.strftime("%d.%m.%Y")
    if isinstance(wert, (Decimal, float)):
        d = Decimal(str(wert))
        if stellen is not None:
            d = d.quantize(Decimal(1).scaleb(-stellen), rounding=ROUND_HALF_UP)
        text = format(d.normalize() if stellen is None else d, "f")
        return text.replace(".", ",")
    return str(wert)


def _csv_bytes(spalten: list[str], zeilen: list[list[Any]], geld: set[str] = frozenset()) -> bytes:
    puffer = io.StringIO()
    w = csv.writer(puffer, delimiter=";", lineterminator="\r\n")
    w.writerow(spalten)
    for z in zeilen:
        w.writerow([_de(v, 2 if s in geld else None) for s, v in zip(spalten, z)])
    return ("\ufeff" + puffer.getvalue()).encode("utf-8")


def _datum_id(d: Optional[date | datetime]) -> Optional[int]:
    return None if d is None else int(d.strftime("%Y%m%d"))


_MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
           "November", "Dezember"]
_TAGE = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]

DAX_MEASURES = """// FlowAudit-AI – DAX-Measures (Tabelle "Measures" anlegen und einfügen)
Rechnungen = COUNTROWS ( fact_rechnung )
Touchless Rechnungen = CALCULATE ( [Rechnungen], dim_workflow_schritt[TOUCHLESS] = 1 )
Dunkelverarbeitungsquote = DIVIDE ( [Touchless Rechnungen], [Rechnungen] )
Maverick-Buying-Quote = DIVIDE ( CALCULATE ( [Rechnungen], fact_rechnung[WORKFLOW_SCHRITT] = "COMPLIANCE_MAVERICK_BUYING" ), [Rechnungen] )
Volumen brutto = SUM ( fact_rechnung[BRUTTO] )
Offenes Volumen = CALCULATE ( [Volumen brutto], ISBLANK ( fact_rechnung[ABSCHLUSS] ) )
Ø Durchlaufzeit (h) = AVERAGE ( fact_rechnung[DURCHLAUFZEIT_H] )
Skonto gefährdet = CALCULATE ( SUM ( fact_rechnung[SKONTO_BETRAG] ), fact_rechnung[SKONTO_STATUS] = "GEFÄHRDET" )
Skonto verloren = CALCULATE ( SUM ( fact_rechnung[SKONTO_BETRAG] ), fact_rechnung[SKONTO_STATUS] = "VERLOREN" )
Skonto gesichert = CALCULATE ( SUM ( fact_rechnung[SKONTO_BETRAG] ), fact_rechnung[SKONTO_STATUS] = "GESICHERT" )
Prüffehler = CALCULATE ( COUNTROWS ( fact_pruefung ), fact_pruefung[STATUS] = "FEHLER" )
Fehler je Rechnung = DIVIDE ( [Prüffehler], [Rechnungen] )
"""

POWER_QUERY = """// FlowAudit-AI – Power Query (M): CSV mit deutscher Kultur laden
// Erweiterter Editor: Pfad anpassen; je Tabelle eine Abfrage (Dateiname austauschen).
let
    Ordner = "C:\\\\FlowAudit\\\\powerbi\\\\",
    Quelle = Csv.Document(File.Contents(Ordner & "fact_rechnung.csv"),
        [Delimiter = ";", Encoding = 65001, QuoteStyle = QuoteStyle.Csv]),
    Kopf = Table.PromoteHeaders(Quelle, [PromoteAllScalars = true]),
    Typen = Table.TransformColumnTypes(Kopf, {
        {"GJAHR", Int64.Type}, {"RECHNUNGSDATUM_ID", Int64.Type}, {"EINGANG_DATUM_ID", Int64.Type},
        {"EINGANG", type datetime}, {"ABSCHLUSS", type datetime}, {"DURCHLAUFZEIT_H", type number},
        {"NETTO", Currency.Type}, {"STEUER", Currency.Type}, {"BRUTTO", Currency.Type},
        {"SKONTO_PROZ", type number}, {"SKONTO_BETRAG", Currency.Type}, {"TOUCHLESS", Int64.Type}
    }, "de-DE")
in
    Typen
"""


def erstelle_powerbi_paket(db_pfad: Path | str, stichtag: Optional[date] = None,
                           daten: Optional[ExportDaten] = None) -> bytes:
    d = daten or lade_exportdaten(db_pfad, stichtag)
    skonto_von = {s["BELNR"]: s for s in d.skonto}
    dateien: dict[str, bytes] = {}
    alle_daten: set[date] = {d.stichtag}

    sp = ["RECHNUNG_ID", "BELNR", "GJAHR", "XBLNR", "LIFNR", "RECHNUNGSDATUM_ID", "EINGANG_DATUM_ID", "EINGANG",
          "ABSCHLUSS", "DURCHLAUFZEIT_H", "NETTO", "STEUER", "BRUTTO", "WORKFLOW_SCHRITT", "URTEIL", "ZAHLSPERRE",
          "FORMAT", "EBELN", "SKONTO_PROZ", "SKONTO_BETRAG", "SKONTO_FRIST_ID", "SKONTO_STATUS", "TOUCHLESS",
          "DATEI_SHA256"]
    zeilen = []
    for r in d.rechnungen:
        eingang, abschluss, bldat = _dt(r["ZZ_EINGANG"]), _dt(r["ZZ_ABSCHLUSS"]), _d(r["BLDAT"])
        s = skonto_von.get(r["BELNR"])
        alle_daten.update(x for x in (bldat, eingang and eingang.date(), s and s["FRIST"]) if x)
        dlz = _r2(Decimal(str((abschluss - eingang).total_seconds() / 3600))) if abschluss and eingang else None
        zeilen.append([f"{r['BELNR']}-{r['GJAHR']}", r["BELNR"], r["GJAHR"], r["XBLNR"], r["LIFNR"], _datum_id(bldat),
                       _datum_id(eingang), eingang, abschluss, dlz, _dez(r["ZZ_NETTO"]), _dez(r["WMWST1"]),
                       _dez(r["RMWWR"]), r["ZZ_WF_SCHRITT"], r["ZZ_URTEIL"], r["ZLSPR"], r["ZZ_FORMAT"], r["ZZ_EBELN"],
                       s["PROZ"] if s else None, s["BETRAG"] if s else None, _datum_id(s["FRIST"]) if s else None,
                       s["STATUS"] if s else None, int(r["ZZ_WF_SCHRITT"] == "TOUCHLESS_FREIGABE"), r["ZZ_SHA256"]])
    dateien["fact_rechnung.csv"] = _csv_bytes(sp, zeilen, {"NETTO", "STEUER", "BRUTTO", "SKONTO_BETRAG"})

    sp = ["PRUEFUNG_ID", "RECHNUNG_ID", "PRUEFCODE", "STATUS", "SOLL", "IST", "ABWEICHUNG", "ABWEICHUNG_PROZ",
          "REFERENZ", "GEPRUEFT_AM", "FAKTEN_SHA256"]
    dateien["fact_pruefung.csv"] = _csv_bytes(sp, [
        [i, f"{p['BELNR']}-{p['GJAHR']}", p["PRUEFCODE"], p["STATUS"], _dez(p["SOLL"]), _dez(p["IST"]),
         _dez(p["ABWEICHUNG"]), _dez(p["ABWEICHUNG_PROZ"]), p["REFERENZ"], _dt(p["GEPRUEFT_AM"]), p["FAKTEN_SHA256"]]
        for i, p in enumerate(d.protokoll, start=1)])

    sp = ["LIFNR", "NAME", "STRASSE", "PLZ", "ORT", "LAND", "UST_ID", "ZAHLUNGSBEDINGUNG"]
    dateien["dim_lieferant.csv"] = _csv_bytes(sp, [[l_["LIFNR"], l_["NAME1"], l_["STRAS"], l_["PSTLZ"], l_["ORT01"],
                                                    l_["LAND1"], l_["STCEG"], l_["ZTERM"]] for l_ in d.lieferanten])

    sp = ["WORKFLOW_SCHRITT", "BEZEICHNUNG", "URTEIL", "TOUCHLESS", "ZUSTAENDIG", "SORTIERUNG"]
    dateien["dim_workflow_schritt.csv"] = _csv_bytes(sp, [
        [s.value, SCHRITT_TEXT[s.value], URTEIL_JE_SCHRITT[s].value, int(s is WorkflowSchritt.TOUCHLESS_FREIGABE),
         EMPFAENGER_JE_SCHRITT[s].value.title() if s in EMPFAENGER_JE_SCHRITT else "Kreditorenbuchhaltung", i]
        for i, s in enumerate(URTEIL_JE_SCHRITT, start=1)])

    sp = ["PRUEFCODE", "BEZEICHNUNG", "KATEGORIE", "BPMN_GATEWAY"]
    dateien["dim_pruefcode.csv"] = _csv_bytes(sp, [[c.value, *PRUEFCODE_INFO[c.value]] for c in PruefCode])

    sp = ["DATUM_ID", "DATUM", "JAHR", "QUARTAL", "MONAT", "MONATSNAME", "JAHR_MONAT", "KW", "WOCHENTAG", "WERKTAG"]
    start, ende = min(alle_daten).replace(day=1), max(alle_daten)
    kalender, tag = [], start
    while tag <= ende:
        kalender.append([_datum_id(tag), tag, tag.year, f"Q{(tag.month - 1) // 3 + 1}", tag.month,
                         _MONATE[tag.month - 1], f"{tag.year}-{tag.month:02d}", tag.isocalendar()[1],
                         _TAGE[tag.weekday()], int(tag.weekday() < 5)])
        tag += timedelta(days=1)
    dateien["dim_datum.csv"] = _csv_bytes(sp, kalender)

    dateien["measures.dax"] = DAX_MEASURES.encode("utf-8")
    dateien["laden.pq"] = POWER_QUERY.encode("utf-8")
    dateien["README.txt"] = f"""FlowAudit-AI – Power-BI-Paket (synthetische Demo-Daten, Stichtag {d.stichtag:%d.%m.%Y})
==================================================================================

Import in Power BI Desktop
  1. Daten abrufen > Text/CSV > Datei wählen
  2. Dateiursprung 65001 (UTF-8), Trennzeichen Semikolon
  3. "Daten transformieren" > Datentyp mit Gebietsschema > Deutsch (Deutschland)
     (oder Skript laden.pq im Erweiterten Editor verwenden)

Sternschema (Beziehungen, Kardinalität n:1, Filterrichtung einfach)
  fact_rechnung[LIFNR]             -> dim_lieferant[LIFNR]
  fact_rechnung[WORKFLOW_SCHRITT]  -> dim_workflow_schritt[WORKFLOW_SCHRITT]
  fact_rechnung[EINGANG_DATUM_ID]  -> dim_datum[DATUM_ID]        (aktiv)
  fact_rechnung[RECHNUNGSDATUM_ID] -> dim_datum[DATUM_ID]        (inaktiv, USERELATIONSHIP)
  fact_pruefung[RECHNUNG_ID]       -> fact_rechnung[RECHNUNG_ID]
  fact_pruefung[PRUEFCODE]         -> dim_pruefcode[PRUEFCODE]

  dim_datum als Datumstabelle markieren (Spalte DATUM).

Measures: siehe measures.dax · Integrität: SHA-256 je Datei in manifest.json
Zahlenformat: Dezimalkomma, kein Tausendertrennzeichen · Datum TT.MM.JJJJ · Bankdaten sind nicht enthalten.
""".replace("\n", "\r\n").encode("utf-8")

    manifest = {
        "erzeugt": datetime.now().isoformat(timespec="seconds"),
        "stichtag": d.stichtag.isoformat(),
        "hinweis": "Synthetische Demo-Daten (FlowAudit-AI)",
        "dateien": {name: {"sha256": hashlib.sha256(inhalt).hexdigest(),
                           "zeilen": inhalt.decode("utf-8-sig").count("\r\n") - 1 if name.endswith(".csv") else None}
                    for name, inhalt in dateien.items()},
    }
    dateien["manifest.json"] = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")

    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for name, inhalt in dateien.items():
            z.writestr(f"flowaudit_powerbi/{name}", inhalt)
    return puffer.getvalue()


# ═════════════════════════════════════════════════════════════════════════════
# 5 · Dateien schreiben
# ═════════════════════════════════════════════════════════════════════════════

def speichere_excel_bericht(db_pfad: Path | str, ziel: Path | str, stichtag: Optional[date] = None) -> Path:
    ziel = Path(ziel)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    ziel.write_bytes(erstelle_excel_bericht(db_pfad, stichtag))
    return ziel


def speichere_powerbi_paket(db_pfad: Path | str, ziel: Path | str, stichtag: Optional[date] = None) -> Path:
    ziel = Path(ziel)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    ziel.write_bytes(erstelle_powerbi_paket(db_pfad, stichtag))
    return ziel


if __name__ == "__main__":  # pragma: no cover
    import argparse

    ap = argparse.ArgumentParser(description="FlowAudit-AI: Excel-Prüfbericht und Power-BI-Paket erzeugen")
    ap.add_argument("--db", default="data/flowaudit_erp.db")
    ap.add_argument("--ziel", default="exports")
    a = ap.parse_args()
    print(speichere_excel_bericht(a.db, Path(a.ziel) / "flowaudit_pruefbericht.xlsx"))
    print(speichere_powerbi_paket(a.db, Path(a.ziel) / "flowaudit_powerbi.zip"))
